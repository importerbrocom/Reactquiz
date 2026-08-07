#!/usr/bin/env python3
"""
Generates the question-import template files.

    python3 docs/templates/build-template.py

Outputs (in this directory):
    questions-import-template.xlsx   <- give this to content authors
    questions-import-template.csv    <- same columns, for tooling/tests

The admin UI serves the .xlsx from "Questions -> Import -> Download template",
so this script is the single source of truth for the template's shape.
Keep it in sync with docs/import-format.md.
"""

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path(__file__).parent

# ---------------------------------------------------------------- columns ----
# (header, width, wrap, required, help text shown in the instructions sheet)
COLUMNS = [
    ("question_no",    9,  False, False, "Your own reference number. Not stored - it only appears in the error report so you can find the row in your sheet."),
    ("question",      58,  True,  True,  "The question text. 10-2000 characters. Line breaks are allowed (Alt+Enter in Excel)."),
    ("option_a",      26,  True,  True,  "Answer choice A. 1-500 characters."),
    ("option_b",      26,  True,  True,  "Answer choice B."),
    ("option_c",      26,  True,  True,  "Answer choice C."),
    ("option_d",      26,  True,  True,  "Answer choice D."),
    ("correct_option", 13, False, True,  "Which choice is correct: A, B, C or D. Lower case and 1/2/3/4 are also accepted."),
    ("explanation",   62,  True,  True,  "Why the correct answer is correct. 5-2000 characters. Shown to the student after a wrong answer, so make it teach something."),
    ("topic",         24,  False, False, "Subject name. Pick from the dropdown (list is on the 'topics' sheet). Strongly recommended - it powers the weak/strong subject analysis on the student's progress page."),
    ("tags",          24,  True,  False, "Optional finer labels, separated by semicolons: autonomic;beta-blockers;high-yield"),
    ("source",        22,  False, False, "Optional. Where the question came from, e.g. 'FMGE 2023 Paper 1' or a textbook reference."),
    ("image_filename",20,  False, False, "Leave blank. Only used if you upload a ZIP of images alongside this file. Images can also be attached later in the admin editor."),
    ("status",        11,  False, False, "active (default) or draft. Use draft for questions that are written but not ready to be served."),
]

HEADERS = [c[0] for c in COLUMNS]

# ----------------------------------------------------------------- topics ----
TOPICS = [
    "Anatomy", "Physiology", "Biochemistry",
    "Pathology", "Pharmacology", "Microbiology", "Forensic Medicine", "Community Medicine",
    "General Medicine", "General Surgery", "Obstetrics & Gynaecology", "Paediatrics",
    "Orthopaedics", "ENT", "Ophthalmology", "Psychiatry", "Dermatology",
    "Anaesthesia", "Radiology",
]

# ------------------------------------------------------------ sample rows ----
SAMPLES = [
    [1, "Which chamber of the heart pumps oxygenated blood into the aorta?",
     "Right atrium", "Right ventricle", "Left atrium", "Left ventricle", "D",
     "The left ventricle has the thickest myocardium because it must generate enough pressure to drive blood through the entire systemic circulation via the aorta. The right ventricle pumps only to the lungs, at roughly one-fifth of that pressure.",
     "Anatomy", "cardiovascular;high-yield", "FMGE 2023 Paper 1", "", "active"],

    [2, "A patient presents with fever, night sweats and weight loss over six weeks. Which investigation is most appropriate first?",
     "Chest X-ray", "MRI brain", "Echocardiogram", "Colonoscopy", "A",
     "This triad strongly suggests pulmonary tuberculosis. A chest X-ray is the standard first-line investigation: it is inexpensive, widely available and highly informative in suspected pulmonary TB.",
     "General Medicine", "tuberculosis;diagnosis", "", "", "active"],

    [3, "What does the abbreviation \"NPO\" mean in a patient's chart?",
     "Nothing by mouth", "No pain observed", "Normal pulse oximetry", "Not previously operated", "A",
     "NPO is from the Latin nil per os, meaning nothing by mouth. It is ordered before surgery and certain procedures to reduce the risk of pulmonary aspiration.",
     "General Surgery", "terminology", "", "", "active"],

    [4, "How many pairs of cranial nerves arise from the brain and brainstem?",
     "10", "11", "12", "13", 3,
     "There are 12 pairs of cranial nerves. Unlike spinal nerves, they emerge directly from the brain and brainstem rather than from the spinal cord.",
     "Anatomy", "neuroanatomy", "", "", "active"],

    [5, "Deficiency of which vitamin causes scurvy?",
     "Vitamin A", "Vitamin C", "Vitamin D", "Vitamin K", "B",
     "Scurvy is caused by vitamin C (ascorbic acid) deficiency. Vitamin C is an essential cofactor for collagen synthesis, so deficiency produces bleeding gums, poor wound healing and perifollicular haemorrhages.",
     "Biochemistry", "vitamins", "", "", "active"],

    [6, "Which ion is primarily responsible for the depolarisation phase of a neuronal action potential?",
     "Sodium (Na+)", "Potassium (K+)", "Chloride (Cl-)", "Calcium (Ca2+)", "A",
     "Voltage-gated sodium channels open first, and rapid Na+ influx down its electrochemical gradient drives the membrane potential sharply positive. Potassium efflux then repolarises the membrane.",
     "Physiology", "neurophysiology", "", "", "active"],

    [7, "Identify the structure labelled X in the accompanying diagram of the nephron.",
     "Glomerulus", "Loop of Henle", "Distal convoluted tubule", "Collecting duct", "B",
     "The structure descends into the renal medulla and forms a hairpin turn, which is characteristic of the loop of Henle. It establishes the medullary osmotic gradient needed to concentrate urine.",
     "Physiology", "renal;diagram", "Physiology Atlas", "nephron-labelled.png", "active"],

    [8, "What is the accepted normal resting heart rate range for a healthy adult?",
     "40-60 bpm", "60-100 bpm", "100-120 bpm", "120-140 bpm", "B",
     "60-100 beats per minute is the accepted normal adult resting range. Below 60 is bradycardia and above 100 is tachycardia, although well-trained athletes are often normally bradycardic.",
     "General Medicine", "vital-signs", "", "", "draft"],
]

# ------------------------------------------------------------------ style ----
HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
REQ_FILL = PatternFill("solid", fgColor="EEF2FF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def build_workbook() -> Workbook:
    wb = Workbook()

    # ---------------------------------------------------- sheet: questions --
    ws = wb.active
    ws.title = "questions"

    ws.append(HEADERS)
    for idx, (name, width, wrap, required, _help) in enumerate(COLUMNS, start=1):
        col = get_column_letter(idx)
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        ws.column_dimensions[col].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for row in SAMPLES:
        ws.append(row)

    for r in range(2, len(SAMPLES) + 2):
        for idx, (_name, _w, wrap, required, _h) in enumerate(COLUMNS, start=1):
            c = ws.cell(row=r, column=idx)
            c.alignment = Alignment(wrap_text=wrap, vertical="top")
            c.border = BORDER
            if required:
                c.fill = REQ_FILL
        ws.row_dimensions[r].height = 58

    # dropdowns, applied generously so they work as rows are added
    last = 3000
    dv_opt = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=False,
                            showErrorMessage=True, errorTitle="Invalid choice",
                            error="Enter A, B, C or D.")
    dv_status = DataValidation(type="list", formula1='"active,draft"', allow_blank=True,
                               showErrorMessage=True, errorTitle="Invalid status",
                               error="Enter active or draft (or leave blank for active).")
    dv_topic = DataValidation(type="list",
                              formula1=f"=topics!$A$2:$A${len(TOPICS) + 1}",
                              allow_blank=True, showErrorMessage=False)
    for dv in (dv_opt, dv_status, dv_topic):
        ws.add_data_validation(dv)
    col_of = {name: get_column_letter(i) for i, (name, *_r) in enumerate(COLUMNS, start=1)}
    dv_opt.add(f"{col_of['correct_option']}2:{col_of['correct_option']}{last}")
    dv_status.add(f"{col_of['status']}2:{col_of['status']}{last}")
    dv_topic.add(f"{col_of['topic']}2:{col_of['topic']}{last}")

    # ------------------------------------------------- sheet: instructions --
    ins = wb.create_sheet("instructions")
    ins.column_dimensions["A"].width = 26
    ins.column_dimensions["B"].width = 104

    def line(a="", b="", bold_a=True, wrap=True):
        ins.append([a, b])
        r = ins.max_row
        if bold_a:
            ins.cell(row=r, column=1).font = BOLD
        ins.cell(row=r, column=1).alignment = Alignment(vertical="top")
        ins.cell(row=r, column=2).alignment = Alignment(wrap_text=wrap, vertical="top")
        return r

    ins.append(["How to fill this file"])
    ins.cell(row=1, column=1).font = TITLE_FONT
    ins.append([])

    line("1. Delete the examples",
         "Rows 2-9 of the 'questions' sheet are examples. Delete them before you import, or set their status to draft.")
    line("2. One row = one question",
         "One row holds the question, its four choices, which choice is correct, and the explanation. Keep each question complete on its own row.")
    line("3. Do not assign days",
         "There is no quiz-day column. Put all 300 questions for a course in this one file; the system gives every student their own personal shuffle of those 300 across the 30 days.")
    line("4. Required columns",
         "question, option_a, option_b, option_c, option_d, correct_option, explanation. A row missing any of these is rejected and reported - it is never imported half-complete.")
    line("5. Explanation is required",
         "Every question needs one. It is what the student reads after answering incorrectly, so it should teach, not just restate the answer.")
    line("6. Fill in topic",
         "Technically optional, but the student's weak/strong subject analysis is empty without it, and days are balanced across subjects using it. Use the dropdown so spellings stay consistent.")
    line("7. NEVER name a letter",
         "Answer choices are reshuffled every time a student sees the question, so that they learn the material instead of memorising a position. This means an explanation must never say 'the answer is B' or 'option C is wrong' - it will be wrong on most students' screens. Always name the answer TEXT instead: 'The left ventricle is correct because...'.")
    line("7b. 'All of the above'",
         "Safe to use - it is automatically pinned to the last position. But avoid 'Both A and B' style choices, because the letters are part of the meaning; write 'Both hypertension and diabetes' instead.")
    line("8. Formatting",
         "Plain text is safest. Bold, italic, underline, superscript, subscript and line breaks are preserved. Pasting from Word or Google Docs is fine - it gets cleaned automatically.")
    line("9. Duplicates",
         "Identical questions are detected automatically (whitespace, capitalisation and choice order are ignored) and reported instead of imported twice. Re-uploading a corrected file is always safe.")
    line("10. How many questions",
         "Exactly 300 per course for a 30-day / 10-per-day programme. The course cannot be published with fewer. You can upload in several files - for example 50 at a time - into the same course.")
    line("11. Images",
         "Leave image_filename blank and attach images afterwards in the admin question editor. Or fill in filenames and upload a ZIP of the images with this file.")
    line("12. Saving",
         "Keep this as .xlsx. If you must use CSV, save as 'CSV UTF-8' so that symbols and accents survive.")
    ins.append([])
    line("Import screen",
         "Upload the file, review the summary (valid / warnings / rejected / duplicates), fix anything flagged, then approve. Nothing is written to the live question bank until you approve it.")

    # -------------------------------------------------------- sheet: topics --
    tp = wb.create_sheet("topics")
    tp.column_dimensions["A"].width = 30
    tp.append(["topic"])
    tp.cell(row=1, column=1).fill = HEADER_FILL
    tp.cell(row=1, column=1).font = HEADER_FONT
    for t in TOPICS:
        tp.append([t])
    tp.append([])
    tp.append(["Add or remove subjects to suit the exam. Keep the spelling identical across every file."])
    tp.cell(row=tp.max_row, column=1).font = Font(italic=True, size=9)

    return wb


def build_csv() -> None:
    path = OUT_DIR / "questions-import-template.csv"
    with path.open("w", newline="", encoding="utf-8-sig") as f:  # BOM: Excel-safe
        w = csv.writer(f, quoting=csv.QUOTE_MINIMAL)
        w.writerow(HEADERS)
        w.writerows(SAMPLES)
    print(f"wrote {path.name}")


if __name__ == "__main__":
    wb = build_workbook()
    xlsx = OUT_DIR / "questions-import-template.xlsx"
    wb.save(xlsx)
    print(f"wrote {xlsx.name}  ({len(HEADERS)} columns, {len(SAMPLES)} sample rows)")
    build_csv()
