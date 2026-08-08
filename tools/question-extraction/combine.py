#!/usr/bin/env python3
"""
Merge per-paper extractions into one import file for a level.

Deduplicates ACROSS papers using the same hash the API enforces
(App\\Models\\Question::makeHash), so a question repeated between years is caught
here rather than being rejected one row at a time at import time.
"""
from __future__ import annotations

import argparse
import hashlib
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

COLUMNS = ["question_no", "question", "option_a", "option_b", "option_c", "option_d",
           "correct_option", "explanation", "topic", "tags", "source",
           "image_filename", "status"]


def question_hash(row: dict) -> str:
    def norm(v: str) -> str:
        return re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", str(v or ""))).strip().lower()
    opts = sorted(norm(row.get(f"option_{k}")) for k in "abcd")
    return hashlib.sha256((norm(row.get("question")) + "|" + "|".join(opts)).encode()).hexdigest()


def read_sheet(path: Path, sheet: str = "questions") -> list[dict]:
    wb = load_workbook(path)
    if sheet not in wb.sheetnames:
        return []
    ws = wb[sheet]
    header = [c.value for c in ws[1]]
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        if record.get("question"):
            out.append({c: record.get(c) or "" for c in COLUMNS})
    return out


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("inputs", nargs="+", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--limit", type=int, help="cap the output at this many rows")
    args = ap.parse_args()

    seen: dict[str, str] = {}
    merged: list[dict] = []
    duplicates: list[tuple[str, str, str]] = []

    for path in args.inputs:
        rows = read_sheet(path)
        kept = 0
        for row in rows:
            h = question_hash(row)
            if h in seen:
                duplicates.append((path.name, str(row["question"])[:70], seen[h]))
                continue
            seen[h] = f"{path.name}:{row['question_no']}"
            row["question_no"] = str(len(merged) + 1)
            merged.append(row)
            kept += 1
        print(f"  {path.name:<28} {len(rows):>4} rows -> {kept:>4} kept")

    if args.limit:
        merged = merged[: args.limit]

    wb = Workbook()
    ws = wb.active
    ws.title = "questions"
    ws.append(COLUMNS)
    widths = [10, 60, 26, 26, 26, 26, 13, 62, 22, 24, 20, 16, 10]
    for i, name in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=i)
        c.fill = PatternFill("solid", fgColor="4F46E5")
        c.font = Font(bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.freeze_panes = "A2"
    for row in merged:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", str(row[c])) for c in COLUMNS])
        for i in (2, 8):
            ws.cell(row=ws.max_row, column=i).alignment = Alignment(wrap_text=True, vertical="top")

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)

    print(f"\n  merged rows        : {len(merged)}")
    print(f"  cross-paper dupes  : {len(duplicates)}")
    for name, q, first in duplicates[:10]:
        print(f"    {name}: {q!r}\n      already seen as {first}")
    with_expl = sum(1 for r in merged if str(r["explanation"]).strip())
    print(f"  with explanations  : {with_expl}/{len(merged)}")
    print(f"\n  wrote {args.out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
